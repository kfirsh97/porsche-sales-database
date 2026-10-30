"""Downloads: Excel with the SQL beside the data, CSV that Excel opens in Hebrew, the whole query not the screen."""
import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import exports  # noqa: E402

COLS = ["אזור", "Sales", "Margin"]
ROWS = [("מרכז", 1200, 0.125), ("צפון", 800, None)]


def test_excel_has_the_data_and_the_query():
    wb = openpyxl.load_workbook(io.BytesIO(exports.to_excel(COLS, ROWS, sql="SELECT 1\nFROM t")))
    assert wb.sheetnames == ["Data", "SQL"]
    data = [[c.value for c in r] for r in wb["Data"].iter_rows()]
    assert data == [COLS, ["מרכז", 1200, 0.125], ["צפון", 800, None]]
    assert wb["Data"].freeze_panes == "A2" and wb["Data"]["A1"].font.bold
    assert [c.value for c in wb["SQL"]["A"]][3:] == ["SELECT 1", "FROM t"]


def test_csv_opens_in_excel_with_hebrew():
    out = exports.to_csv(COLS, ROWS)
    assert out.startswith(b"\xef\xbb\xbf")                  # the byte-order mark Excel needs for UTF-8
    assert out.decode("utf-8-sig").splitlines()[1] == "מרכז,1200,0.125"


def test_file_names_are_safe():
    assert exports.file_name("Sales by region?", "xlsx") == "sales-by-region.xlsx"
    assert exports.file_name("מכירות לפי אזור", "csv") == "מכירות-לפי-אזור.csv"
    assert exports.file_name("", "csv") == "result.csv"


def test_full_result_goes_through_the_read_only_check():
    with pytest.raises(ValueError, match="Only SELECT"):
        exports.full_result("DELETE FROM sales")
