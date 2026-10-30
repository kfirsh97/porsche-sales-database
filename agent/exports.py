"""Results out of the app: every answer and dashboard card downloads as Excel or CSV.

The file is made from the whole query, run again when the button is pressed - not from the first rows on screen.
Excel gets a second sheet with the SQL, so whoever opens the file can see (and rerun) where the numbers came from.
CSV is written with a byte-order mark, so Excel opens Hebrew text in it correctly.
"""
import io
import re
from datetime import datetime

import pandas as pd

MAX_EXPORT_ROWS = 200_000          # Excel's own limit is ~1M rows; past this a database export tool fits better


def frame(columns, rows) -> pd.DataFrame:
    return pd.DataFrame([list(r) for r in rows], columns=columns)


def full_result(sql: str):
    """The whole result of a query, through the same read-only checks as the agent's own queries."""
    from tools import run_query_all
    return run_query_all(sql, max_rows=MAX_EXPORT_ROWS)


def to_csv(columns, rows) -> bytes:
    return frame(columns, rows).to_csv(index=False).encode("utf-8-sig")


def to_excel(columns, rows, sql: str = "", title: str = "") -> bytes:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    df = frame(columns, rows)
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name="Data", index=False)
        ws = xl.sheets["Data"]
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, 1):
            ws.cell(row=1, column=i).font = Font(bold=True)
            sample = [str(col)] + [str(v) for v in df[col].head(200).tolist()]
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(s) for s in sample) + 2, 50)
            if pd.api.types.is_float_dtype(df[col]):
                for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell[0].number_format = "#,##0.00"
        if sql:
            info = xl.book.create_sheet("SQL")
            info["A1"], info["A2"] = "Query", f"Exported {datetime.now():%d %b %Y, %H:%M} - {len(df):,} rows"
            info["A1"].font = Font(bold=True)
            for n, line in enumerate(sql.strip().splitlines(), 4):
                info.cell(row=n, column=1, value=line)
            info.column_dimensions["A"].width = 100
    return out.getvalue()


def file_name(title: str, ext: str) -> str:
    stem = re.sub(r"[^\w֐-׿]+", "-", (title or "result").strip()).strip("-").lower()[:60] or "result"
    return f"{stem}.{ext}"
